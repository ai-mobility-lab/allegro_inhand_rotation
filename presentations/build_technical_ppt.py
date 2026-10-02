"""Build an editable, source-grounded three-slide research overview."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = Path(__file__).resolve().parent
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
NAVY, TEAL, GRAY, LIGHT, WHITE, AMBER = '132B43', '007F82', '536578', 'F0F5F8', 'FFFFFF', 'B86A13'

def box(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    o = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    o.fill.solid(); o.fill.fore_color.rgb = RGBColor.from_string(fill)
    o.line.fill.background() if line is None else None
    if line: o.line.color.rgb = RGBColor.from_string(line)
    return o

def txt(s, x, y, w, h, text, size=17, color=NAVY, bold=False, font='Liberation Sans'):
    o = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = o.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.015)
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = font; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(8)
    return o

def slide(n, title, subtitle, sources):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, 13.333, .09, TEAL)
    txt(s, .5, .28, 12, .25, 'ALLEGRO IN-HAND ROTATION   /   TECHNICAL REPOSITORY COMPARISON', 10, TEAL, True)
    txt(s, .5, .76, 12.3, .55, title, 28, NAVY, True)
    txt(s, .5, 1.40, 12.2, .48, subtitle, 16, GRAY)
    box(s, .5, 7.00, 12.3, .015, 'D9E2E9')
    txt(s, .5, 7.11, 11.9, .23, sources, 9, GRAY)
    txt(s, 12.25, 7.07, .6, .3, f'{n:02d}', 14, TEAL, True)
    return s

def arrow(s, x, y, w=.38):
    box(s, x, y, w, .22, TEAL, shape=MSO_SHAPE.CHEVRON)

def card(s, x, y, w, h, label, body, size=16):
    box(s, x, y, w, h, LIGHT)
    box(s, x, y, .045, h, TEAL)
    txt(s, x+.18, y+.16, w-.36, .4, label, 16, TEAL, True)
    txt(s, x+.18, y+.72, w-.36, h-.81, body, min(size, 14.5))

s = slide(1, 'Isaac Lab migration with the HORA learning interface',
          'The simulator, control plumbing and deployment ordering change; PPO and adaptation source files remain identical.',
          '[1] isaaclab/allegro_hand_hora_env.py + wrapper.py   [2] deploy_ros2*.py   •   Full paths and scope in speaker notes')
for x, w, label, body, fill, col in [
    (.5, 3.1, 'ORIGINAL  /  ISAAC GYM', 'VecTask + gymtorch\nTask + simulation loop', LIGHT, NAVY),
    (4.05, 4.15, 'CURRENT  /  ISAAC LAB', 'DirectRLEnv + Articulation / RigidObject\nAction / observation / reset hooks', NAVY, WHITE),
    (8.65, 4.15, 'HORA COMPATIBILITY WRAPPER', 'obs · priv_info · proprio_hist\nterminated ∨ truncated → done', LIGHT, NAVY)]:
    box(s,x,2.1,w,1.28,fill)
    txt(s,x+.16,2.24,w-.32,.26,label,12,TEAL if fill==LIGHT else '66DDD5',True)
    txt(s,x+.16,2.67,w-.32,.65,body,15,col)
arrow(s,3.66,2.62); arrow(s,8.24,2.62)
card(s,.5,3.67,3.95,2.48,'01  Policy contract',
     '16 actions; 96-D observation\n3 × [16 joint positions + 16 targets]\n30 × 32 proprioceptive history\nStage 1 teacher → Stage 2 adaptation',16)
card(s,4.67,3.67,3.95,2.48,'02  Control semantics',
     '120 Hz physics / 20 Hz policy\nq* ← clip(q* + a / 24)\nτ = clip(Kp(q* − q) − Kd q̇, ±0.5)\nq̇ uses the previous physics state.',16)
card(s,8.84,3.67,3.96,2.48,'03  Deployment alignment',
     'Finger-major ↔ joint-rank-major\n4 × 4 transpose permutation aligns observations, history and actions.\nTimeout flags retained for PPO.',16)
txt(s,.53,6.42,12,.35,'Research implication: preserve the learning design while rebuilding simulation and hardware interfaces.',17,TEAL,True)
s.notes_slide.notes_text_frame.text = '''Comparison scope: local original /home/changsik/projects/inhand_manipulation/isaacgym_repo/allegro_inhand_rotation (HEAD 52d076c plus local edits to train_s1.sh and vis_s1.sh) versus /home/changsik/projects/inhand_manipulation/allegro_inhand_rotation (HEAD 79656ee plus all working-tree changes, including untracked export code), inspected 2026-09-21. This is an implementation comparison, not an experimental benchmark.
Source [1]: hora/tasks/isaaclab/allegro_hand_hora_env.py:284,703,726,761; hora/tasks/isaaclab/wrapper.py; configs/task/AllegroHandHora.yaml:9,20,39,98. Original: hora/tasks/base/vec_task.py and hora/tasks/allegro_hand_hora.py.
Byte comparison found hora/algo/ppo/*.py, hora/algo/padapt/*.py and hora/algo/models/*.py unchanged relative to the supplied original. Same source does not establish numerical or checkpoint equivalence across simulators.
The wrapper preserves time_outs separately from combined done flags. Observations contain normalized joint positions and joint targets, not measured joint velocities. The 30-step history is 30 x 32.
Control equation describes inherited incremental target / torque semantics ported to Isaac Lab; it is not proposed as a new controller. The new _apply_action stores previous joint positions to avoid a zero finite difference from lazy state reads.
Source [2]: hora/algo/deploy/deploy_ros2.py:63 and deploy_ros2_two_hands.py. The explicit self-inverse permutation is [0,4,8,12,1,5,9,13,2,6,10,14,3,7,11,15]. This aligns new simulation order with existing HORA/ROS mappings. No hardware validation was performed for this presentation.'''

s = slide(2, 'DIGIT sensing and mesh-object simulation',
          'A left-hand DIGIT variant adds compliant contact surfaces, calibrated tactile rendering and mesh-object support.',
          '[3] LeftAllegroHandDigitHora.yaml + allegro_hand_grasp_env.py   [4] dataset_collection/sensors.py + object spawning')
txt(s,.5,2.06,5.95,.35,'WHAT CHANGED RELATIVE TO THE ORIGINAL',13,TEAL,True)
rows=[('Hand geometry','Standard HORA / V4 assets','Left DIGIT housings + 4 elastomer pads'),
      ('Grasp detection','CPU get_env_rigid_contacts','GPU ContactSensor per fingertip → object'),
      ('Object geometry','Primitive URDF objects','Native primitives + YCB / DexTouch meshes')]
for i,(label,old,new) in enumerate(rows):
    y=2.59+i*1.08
    box(s,.5,y,5.95,.94,LIGHT)
    txt(s,.66,y+.09,5.6,.22,label.upper(),11,TEAL,True)
    txt(s,.66,y+.36,5.62,.25,old,14,GRAY)
    txt(s,.66,y+.64,5.62,.27,'→ '+new,15,NAVY,True)
txt(s,.58,6.01,5.8,.65,'Stable grasps: ≥2 contacts above 0.1 N,\nplus proximity and height checks after warm-up.',16,NAVY)
box(s,6.78,2.08,6.02,4.65,NAVY)
txt(s,7.00,2.27,5.5,.35,'TACTILE IMAGE GENERATION',13,'66DDD5',True)
for y,label,body in [(2.85,'01  Contact geometry','Rigid elastomer links use a compliant PhysX material.'),
                     (3.9,'02  Optical measurement','4 fingertip cameras + a pre-reset, no-contact baseline.'),
                     (4.95,'03  Tactile RGB + contact mask','Calibrated Taxim rendering via VisuoTactileSensor.')]:
    txt(s,7.03,y,5.45,.3,label,18,WHITE,True)
    txt(s,7.03,y+.39,5.4,.5,body,16,'D5E2EC')
txt(s,7.03,6.05,5.45,.38,'Contact pixel: d₀ − d > ε    (default ε = 0.3 mm)',16,'66DDD5',True)
s.notes_slide.notes_text_frame.text = '''Source [3]: configs/task/LeftAllegroHandDigitHora.yaml and LeftAllegroHandDigitGrasp.yaml; assets/allegro/allegro_digit_left_elastomer.urdf. The original already supports standard left/right hands and ROS deployment; these are inherited features, not new contributions.
hora/tasks/isaaclab/allegro_hand_grasp_env.py:90-141 replaces the original CPU contact query with one object-filtered ContactSensor per fingertip. The code comments explain that combining multiple fingertip bodies into one filtered sensor produced zero forces in vectorized environments. After a 10-step warm-up, grasp validity requires all tips within 0.1 m of the object, at least two force magnitudes >0.1 N, and object height above threshold. Successful full episodes populate the grasp cache. GPU-compatible detection is a code capability; no throughput benchmark is claimed.
Source [4]: hora/tasks/isaaclab/allegro_hand_hora_env.py:102-184,314-323,563-685 for compliant materials and mesh spawning; scripts/dataset_collection/sensors.py:70-124,309-432 for calibrated DIGIT cameras and baseline capture. Primitive shapes use native Isaac Lab spawn configurations. Meshes use MeshConverter and cached USD with per-environment object/scale selection.
The elastomer is a rigid link with a compliant contact material, not a finite-element gel simulation. VisuoTactileSensor enables camera tactile output and disables force-field output. Calibration assets are required for Taxim rendering.
scripts/dataset_collection/neuralfeels_export.py:64 implements contact = finite(depth) AND depth>0 AND nominal-depth>mask_eps. Collector default mask_eps is 3e-4 m. No contact pixels are fabricated. Tactile sensing is added to dataset collection; the unchanged HORA policy remains proprioceptive rather than tactile-conditioned.'''

s = slide(3, 'From policy rollouts to NeuralFeels datasets',
          'New Stage 1 and Stage 2 collectors export synchronized vision, touch, kinematics and object ground truth.',
          '[5] collect_stage{1,2}_feelsight_dataset.py   [6] dataset_collection/{neuralfeels_export,legacy_tactile,feelsight_writer,gt_sdf}.py')
nodes=[(.5,2.67,'POLICY ROLLOUT','Teacher or adapted student'),(3.58,2.67,'SYNCHRONIZED CAPTURE','4 DIGIT + scene RGB-D'),(6.66,2.67,'FEELSIGHT EXPORT','Images + state + SDF'),(9.74,3.06,'NEURALFEELS INPUT','Existing GT-depth loader')]
for x,w,label,body in nodes:
    box(s,x,2.13,w,1.02,NAVY)
    txt(s,x+.13,2.30,w-.26,.25,label,11,'66DDD5',True)
    txt(s,x+.13,2.71,w-.26,.3,body,14,WHITE)
for x in [3.24,6.32,9.4]: arrow(s,x,2.53,.27)
card(s,.5,3.47,3.95,2.63,'01  Frame consistency',
     'Map joints by name; swap index / ring labels for NeuralFeels.\nTworld,cam = Tworld,tip · Ttip,cam\nEvery saved camera pose is checked against NeuralFeels FK.',16)
card(s,4.67,3.47,3.95,2.63,'02  Metric depth encoding',
     'OpenGL camera axes; visible depth has negative camera-z.\nz = uint8 / scale + offset\nFit scale / offset per episode; preserve masks with lossless storage.',16)
card(s,8.84,3.47,3.96,2.63,'03  Timing + ground truth',
     'Render every physics step; exclude terminal auto-reset frames.\nSave calibrated K, RGB-D, masks, hand / object poses and joint states.\nExport object SDF at 0.5 mm pitch.',16)
box(s,.5,6.32,12.3,.47,'E3F2EF')
txt(s,.65,6.40,12,.29,'Research capability: calibrated tactile reconstruction datasets. Performance gains require experimental validation.',14,TEAL,True)
s.notes_slide.notes_text_frame.text = '''Source [5]: scripts/collect_stage1_feelsight_dataset.py:199,251,299-315 and scripts/collect_stage2_feelsight_dataset.py. Both collectors are new relative to the original repository. Rendering interval is set to 1. DirectRLEnv resets before returning done, so the terminal step is excluded to avoid mixing reset-state imagery into the prior episode. Nonempty short episodes are saved.
Source [6]: scripts/dataset_collection/neuralfeels_export.py, conventions.py, legacy_tactile.py, feelsight_writer.py, gt_sdf.py and README.md. The export targets existing NeuralFeels loaders in GT tactile-depth mode. Full end-to-end simulator/NeuralFeels execution was not performed for this presentation.
Export joint numbers: [8,9,10,11,4,5,6,7,0,1,2,3,12,13,14,15]. Index and ring sensor names are swapped. Actual sensor world camera matrices are compared with FK using np.allclose(atol=2e-5, rtol=0); this is a matrix-element tolerance, not an angular error bound. Camera offset relative to housing tip is (-0.002132,0.000021,0.017545) m with quaternion wxyz=(0.5,0.5,-0.5,-0.5).
Intrinsics come from actual sensors. Resizing uses fx'=sx*fx, fy'=sy*fy, cx'=(cx+0.5)*sx-0.5, cy'=(cy+0.5)*sy-0.5. Native tactile render is 640x480; default exported tactile image is 240x320 (W x H). Scene default is 640x480.
Tactile depth uses z=depth_uint8/digit_info.depth_scale + digit_info.cam_dist with zero reserved for no contact. The scale and offset are fitted to an episode to avoid saturation. Depth/mask names retain .jpg for existing loaders but payloads are lossless PNG; RGB uses JPEG. Scene metric depth is stored in depth.npz with scale 1. Object ground truth SDF default voxel pitch is 5e-4 m.
The new CPU regression suite tests compatibility against extracted original consumer source without optional SLAM imports; no test result is asserted in this deck. The scope includes current uncommitted/untracked edits as of 2026-09-21. There are no measured performance comparisons in this source review.'''

prs.core_properties.title = 'Allegro In-Hand Rotation: Technical Changes from the Original Repository'
prs.core_properties.subject = 'Isaac Lab migration, DIGIT sensing, and NeuralFeels dataset export'
prs.core_properties.author = 'Research Engineering'
prs.core_properties.keywords = 'Allegro, Isaac Lab, HORA, DIGIT, NeuralFeels, repository comparison'
prs.save(OUT / 'allegro_research_changes.pptx')
print(OUT / 'allegro_research_changes.pptx')
