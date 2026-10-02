"""Build an editable, source-grounded three-slide research overview."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = Path(__file__).resolve().parent
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
NAVY, TEAL, GRAY, LIGHT, WHITE, AMBER = '132B43', '007F82', '536578', 'F0F5F8', 'FFFFFF', 'B86A13'

def box(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    o = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    o.fill.solid(); o.fill.fore_color.rgb = RGBColor.from_string(fill)
    o.line.fill.background() if line is None else None
    if line: o.line.color.rgb = RGBColor.from_string(line)
    return o

def txt(s, x, y, w, h, text, size=17, color=NAVY, bold=False, font='Liberation Sans'):
    o = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = o.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.015)
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = font; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(3)
    return o

def slide(n, title, subtitle, sources):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, 13.333, .09, TEAL)
    txt(s, .5, .28, 12, .25, 'ALLEGRO IN-HAND ROTATION   /   ISAAC GYM → ISAAC LAB / MIGRATION STUDY', 10, TEAL, True)
    txt(s, .5, .76, 12.3, .55, title, 28, NAVY, True)
    txt(s, .5, 1.40, 12.2, .48, subtitle, 16, GRAY)
    box(s, .5, 7.00, 12.3, .015, 'D9E2E9')
    txt(s, .5, 7.11, 11.9, .23, sources, 9, GRAY)
    txt(s, 12.25, 7.07, .6, .3, f'{n:02d}', 14, TEAL, True)
    return s

def arrow(s, x, y, w=.38):
    box(s, x, y, w, .22, TEAL, shape=MSO_SHAPE.CHEVRON)

def card(s, x, y, w, h, label, body, size=16):
    box(s, x, y, w, h, LIGHT)
    box(s, x, y, .045, h, TEAL)
    txt(s, x+.18, y+.16, w-.36, .4, label, 16, TEAL, True)
    txt(s, x+.18, y+.72, w-.36, h-.81, body, min(size, 14.5))

import json
notes = json.loads((OUT / 'source_notes.json').read_text())
official = 'https://isaac-sim.github.io/IsaacLab/main/source/migration/migrating_from_isaacgymenvs.html'

def takeaway(s, text):
    box(s,.5,6.34,12.3,.47,'E3F2EF')
    txt(s,.65,6.43,12,.3,text,15,TEAL,True)

s = slide(1, 'Isaac Gym vs. Isaac Lab: what actually changes?',
    'Both support GPU-parallel robot simulation. The main difference here is how the experiment is built and controlled.',
    'Scope: the original Isaac Gym implementation versus this repository’s Isaac Lab DirectRLEnv port. Sources in speaker notes.')
for x,w,label,body in [
    (.5,5.9,'ORIGINAL: ISAAC GYM','HORA task + custom simulation loop\n↓\nIsaac Gym API → physics + robot-state arrays'),
    (6.9,5.9,'CURRENT: ISAAC LAB','HORA task inside a reusable environment framework\n↓\nIsaac Sim → scene assets + physics + sensors')]:
    box(s,x,2.04,w,1.55,NAVY)
    txt(s,x+.18,2.21,w-.36,.3,label,13,'66DDD5',True)
    txt(s,x+.18,2.65,w-.36,.85,body,16,WHITE)
arrow(s,6.52,2.72,.25)
rows=[
    ('Who runs each step?', 'Repository code advances physics and refreshes state.', 'The framework runs the loop; task functions supply behavior.'),
    ('How is the world described?', 'Create simulation actors and operate on state arrays.', 'Configure robot, object and sensor components in a scene.'),
    ('What does it enable here?', 'A working platform for learning in-hand rotation.', 'The same learning design, plus integrated sensing and data capture.')]
for i,(label,old,new) in enumerate(rows):
    y=3.85+i*.76
    txt(s,.55,y,2.5,.6,label,15,TEAL,True)
    txt(s,3.03,y,4.25,.64,old,16,GRAY)
    txt(s,7.61,y,5.08,.64,new,16,NAVY)
takeaway(s, 'The learning method is retained; the surrounding simulation architecture is rebuilt.')
s.notes_slide.notes_text_frame.text = '''Audience framing: Isaac Gym is the original GPU-oriented simulation platform used by this repository. Isaac Lab is a robot-learning framework built on Isaac Sim. The current repository uses the DirectRLEnv workflow; it does not use a manager-based task design. The original also has a reusable VecTask base class. The distinction on this slide concerns the ownership of the simulation loop and scene abstractions in the compared implementations, not a claim that Isaac Gym had no framework or sensors.
Both systems support GPU physics and parallel environments. This migration is not a CPU-to-GPU conversion. Isaac Gym also has camera capabilities; integrated tactile collection is a new repository feature, not proof that the original platform could never support sensing.
Local evidence: original hora/tasks/base/vec_task.py:step explicitly calls pre_physics_step, fetch_results, simulate and post_physics_step. Current hora/tasks/isaaclab/allegro_hand_hora_env.py subclasses DirectRLEnv and configures Articulation, RigidObject and InteractiveScene. train.py launches Isaac Sim through AppLauncher before importing simulation-dependent modules.
Official migration reference (accessed 2026-09-21): ''' + official + '\nThe implementation-specific conclusions are grounded in local code. Documentation may describe newer framework releases; this deck describes the checked-out port.\n' + notes[0]

s = slide(2, 'The migration changes data flow, timing and configuration',
    'The same robot-learning experiment must be expressed through a different software interface.',
    'Local sources: train.py; hora/tasks/isaaclab/{allegro_hand_hora_env,sim_cfg,wrapper}.py; original tasks/base/vec_task.py.')
for x,w,label in [(.5,2.16,'CONCEPT'),(2.8,4.62,'ISAAC GYM: ORIGINAL'),(7.56,5.24,'ISAAC LAB: CURRENT PORT')]:
    box(s,x,2.05,w,.45,NAVY)
    txt(s,x+.14,2.16,w-.28,.23,label,12,'66DDD5',True)
rows=[
    ('Control timing','Custom loop repeats physics updates for each action.','Framework repeats the action hook: 6 physics steps per decision.'),
    ('Robot measurements','Explicitly refresh shared state arrays before reading them.','Read robot/object data through components; some values update on demand.'),
    ('Physics settings','Settings collected in the original simulation configuration.','Translate settings across the scene, individual robots and objects.'),
    ('Learning interface','Learner expects observations, reward, done and extra data.','A compatibility adapter restores that format and preserves timeout information.')]
for i,(a,b,c) in enumerate(rows):
    y=2.68+i*.83
    box(s,.5,y,12.3,.75,LIGHT if i%2==0 else 'FAFCFD')
    txt(s,.65,y+.13,1.9,.57,a,16,TEAL,True)
    txt(s,2.95,y+.1,4.31,.65,b,16,GRAY)
    txt(s,7.71,y+.1,4.93,.65,c,16)
takeaway(s, 'Preserved interface: 16 finger commands, the same observation format, and unchanged learning code.')
s.notes_slide.notes_text_frame.text = '''Explain the timing in plain language: The learned controller chooses a command 20 times per second. The physics model updates 120 times per second, applying that command across six updates. These are the current configured rates, not a speed benchmark. The original controlFrequencyInv maps to DirectRLEnv decimation. Isaac Gym substeps and controlFrequencyInv are different concepts; do not equate simulator substeps with action decimation.
State access: Original code uses acquired/refreshed state tensors. Current code reads Articulation.data and RigidObject.data. These remain tensor-backed, GPU-compatible mechanisms. The difference is the access and update contract, not arrays disappearing.
Physics settings: sim_cfg.py translates scene-wide values into SimulationCfg and per-asset collision/rigid-body/articulation properties. It drops original knobs without a corresponding mapping. Its custom center-of-mass randomizer also handles the different tensor shape of a single rigid object. Hydra configuration is retained, with nested containers sanitized before passing into Isaac Lab configclasses.
Interface: HoraDirectEnvWrapper combines terminated and truncated for legacy done, while retaining time_outs separately for the learner. It returns obs, priv_info and proprio_hist. Observation size is 96, with 30 x 32 history and 16 actions. The learning, adaptation and model source files match the original.
Launch lifecycle: Isaac Sim AppLauncher must start before imports that touch Isaac Lab. The original train.py did not require that application lifecycle. This is another migration change beyond function renaming.
Equivalence limit: matching interface shapes does not establish identical physics or ready-to-use original checkpoints. The code explicitly notes that friction randomization changes from a shared hand/object sample to independently sampled asset values. These details can change the learning distribution.
Official migration reference: ''' + official + '\n' + notes[0]

s = slide(3, 'Why the port was difficult: hidden assumptions changed',
    'Documented migration issues show why identical-looking code can produce different motion, contacts or images.',
    'Evidence: comments and fixes in allegro_hand_hora_env.py, allegro_hand_grasp_env.py, deploy_ros2*.py and dataset collectors.')
for x,w,label in [(.5,2.75,'CHANGED ASSUMPTION'),(3.39,4.42,'DEVELOPMENT DIFFICULTY'),(7.95,4.85,'IMPLEMENTED FIX')]:
    box(s,x,2.05,w,.45,NAVY)
    txt(s,x+.14,2.16,w-.28,.23,label,12,'66DDD5',True)
rows=[
    ('Joint order','The imported hand lists joints in a different order; commands need remapping.','Reorder measurements and actions so each value belongs to the correct joint.'),
    ('State-update timing','Two joint-position reads can return the same state, giving a zero velocity estimate.','Save the previous physics-step position explicitly before estimating velocity.'),
    ('Contact reporting','A combined fingertip sensor returned zero filtered forces in parallel environments.','Use one object-filtered contact sensor per fingertip; this also supports GPU grasp checks.'),
    ('Asset conversion','The robot-description import path could produce invisible object geometry.','Create simple shapes directly; use a separate mesh-to-scene conversion path.')]
for i,(a,b,c) in enumerate(rows):
    y=2.68+i*.83
    box(s,.5,y,12.3,.75,LIGHT if i%2==0 else 'FAFCFD')
    txt(s,.65,y+.15,2.43,.5,a,16,AMBER,True)
    txt(s,3.54,y+.1,4.1,.65,b,15,GRAY)
    txt(s,8.10,y+.1,4.53,.65,c,15)
takeaway(s, 'Migration preserves the task’s intent; numerical equivalence and performance still need validation.')
s.notes_slide.notes_text_frame.text = '''The table describes code-documented problems in this development environment, not universal defects in Isaac Lab. No claimed development durations, measured speedups or accuracy gains are inferred.
1. Joint ordering: deploy_ros2.py and deploy_ros2_two_hands.py document that the Isaac Lab URDF-to-USD importer changes this asset from finger-major to joint-rank-major order. A self-inverse 4x4 transpose permutation aligns observations, target histories and actions. Existing grasp caches and checkpoints should not be assumed interchangeable across layouts.
2. State timing: allegro_hand_hora_env.py:_apply_action documents that repeated lazy joint-position reads return the same state. _previous_dof_pos stores the prior physics-step positions for finite-difference velocity used by the damping term. The inherited incremental-target and PD control design is retained.
3. Contact: allegro_hand_grasp_env.py:_setup_scene documents all-zero filtered force matrices for a single sensor spanning four fingertips when running multiple environments. One sensor per fingertip, filtered against the object, addresses the indexing limitation. The original get_env_rigid_contacts implementation was CPU-only for grasp generation; this does not imply original policy training was CPU-only.
4. Geometry: allegro_hand_hora_env.py:_build_object_cfg documents an empty/invisible USD payload/reference composition path for imported rigid objects. Native primitive spawning avoids this for simple shapes; MeshConverter provides a separate USD path for detailed meshes. USD is the scene-description format used in Isaac Sim. The code also repairs transparent imported materials.
Additional timing issue for the new dataset collection: DirectRLEnv has already reset by the time done is returned. Collectors exclude that returned frame to avoid recording a new grasp in the previous episode, and render every physics step for synchronized sensor data.
Tradeoffs: the port disables replicate_physics because geometry and scale vary across environments; the code notes a performance cost. No magnitude has been measured for this presentation. No claim is made that Isaac Lab is always faster, more physically accurate or numerically identical to Isaac Gym.
''' + notes[0] + '\n' + notes[1]

prs.core_properties.title = 'Isaac Gym to Isaac Lab: Differences, Migration and Development Challenges'
prs.core_properties.subject = 'Repository-grounded comparison for a multidisciplinary research audience'
prs.core_properties.author = 'Research Engineering'
prs.save(OUT / 'allegro_research_changes.pptx')
print(OUT / 'allegro_research_changes.pptx')
