"""Build an editable, source-grounded three-slide research overview."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUT = Path(__file__).resolve().parent
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
NAVY, TEAL, GRAY, LIGHT, WHITE, AMBER = '132B43', '007F82', '536578', 'F0F5F8', 'FFFFFF', 'B86A13'

def box(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    o = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    o.fill.solid(); o.fill.fore_color.rgb = RGBColor.from_string(fill)
    o.line.fill.background() if line is None else None
    if line: o.line.color.rgb = RGBColor.from_string(line)
    return o

def txt(s, x, y, w, h, text, size=17, color=NAVY, bold=False, font='Liberation Sans'):
    o = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = o.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.015)
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line; p.font.name = font; p.font.size = Pt(size)
        p.font.bold = bold; p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(3)
    return o

def slide(n, title, subtitle, sources):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, 13.333, .09, TEAL)
    txt(s, .5, .28, 12, .25, 'ALLEGRO IN-HAND ROTATION   /   WHAT CHANGED • WHY IT MATTERS • DEVELOPMENT CHALLENGES', 10, TEAL, True)
    txt(s, .5, .76, 12.3, .55, title, 28, NAVY, True)
    txt(s, .5, 1.40, 12.2, .48, subtitle, 16, GRAY)
    box(s, .5, 7.00, 12.3, .015, 'D9E2E9')
    txt(s, .5, 7.11, 11.9, .23, sources, 9, GRAY)
    txt(s, 12.25, 7.07, .6, .3, f'{n:02d}', 14, TEAL, True)
    return s

def arrow(s, x, y, w=.38):
    box(s, x, y, w, .22, TEAL, shape=MSO_SHAPE.CHEVRON)

def card(s, x, y, w, h, label, body, size=16):
    box(s, x, y, w, h, LIGHT)
    box(s, x, y, .045, h, TEAL)
    txt(s, x+.18, y+.16, w-.36, .4, label, 16, TEAL, True)
    txt(s, x+.18, y+.72, w-.36, h-.81, body, min(size, 14.5))

import json
notes = json.loads((OUT / 'source_notes.json').read_text())

def takeaway(s, text):
    box(s, .5, 6.34, 12.3, .47, 'E3F2EF')
    txt(s, .65, 6.43, 12, .3, text, 15, TEAL, True)

s = slide(1, 'From learning to rotate objects to studying touch',
    'The original learns how a robot hand rotates an object. The extension also records what the hand sees and feels.',
    'Comparison of the supplied local repositories, including working-tree edits. Code references and technical details in speaker notes.')
txt(s, .65, 2.07, 2.0, .3, 'AREA', 12, GRAY, True)
box(s, 2.8, 2.0, 4.15, .47, LIGHT)
box(s, 7.12, 2.0, 5.68, .47, NAVY)
txt(s, 2.98, 2.11, 3.8, .25, 'ORIGINAL REPOSITORY', 12, GRAY, True)
txt(s, 7.3, 2.11, 5.2, .25, 'CURRENT REPOSITORY — ADDED / CHANGED', 12, '66DDD5', True)
rows = [
    ('Simulation engine', 'Isaac Gym runs the robot physics.', 'Moved to Isaac Lab; adapted control and robot interfaces.'),
    ('Fingertips', 'Standard robot-hand fingertips.', 'Added four DIGIT touch sensors: small cameras behind soft pads.'),
    ('Object shapes', 'Simple shapes, such as spheres\nand cylinders.', 'Added mesh models of everyday objects\nfrom YCB and DexTouch collections.'),
    ('Research output', 'Learned rotation behavior\nand robot deployment.', 'Also exports touch images, scene images,\n3D distances and known object geometry.'),
]
for i, (area, old, new) in enumerate(rows):
    y = 2.64 + .83*i
    box(s, .5, y, 12.3, .72, LIGHT if i%2 == 0 else 'FAFCFD')
    txt(s, .65, y+.15, 2.05, .5, area, 16, NAVY, True)
    txt(s, 2.98, y+.12, 3.83, .59, old, 16, GRAY)
    txt(s, 7.3, y+.12, 5.24, .59, new, 16, NAVY)
takeaway(s, 'Core learning method retained: the main contribution is simulation, sensing and reusable research data.')
s.notes_slide.notes_text_frame.text = '''Presenter framing: A robot hand must coordinate its fingers to turn an object without dropping it. The original already learns this skill and supports standard left/right hands and ROS deployment. The extension broadens the experimental platform: a new simulator, touch-enabled fingertips, richer object shapes and a dataset pipeline. Do not present the inherited learning algorithm or ROS capability as new work.
Policy means the learned rule that converts recent robot measurements into finger commands. DIGIT is an optical tactile sensor: a camera observes a soft contact surface. A mesh describes a 3D object surface with small polygons. YCB and DexTouch provide everyday-object models.
The learning, adaptation and neural-network source files match the original, but identical source does not establish identical simulator behavior or direct checkpoint compatibility.
''' + notes[0] + '\n' + notes[1]

s = slide(2, 'The new capability: paired “see + touch” experiments',
    'A learned controller moves the object while the simulator records measurements and the true state at the same time.',
    'New collectors: collect_stage1_feelsight_dataset.py / collect_stage2_feelsight_dataset.py and scripts/dataset_collection/.')
steps = [(.5, '1  MOVE', 'The learned controller\nrotates an object.'),
         (3.62, '2  MEASURE', 'Four fingertip cameras\n+ scene cameras.'),
         (6.74, '3  ALIGN', 'Match images to finger\npositions and time.'),
         (9.86, '4  EXPORT', 'Load the recordings\ninto NeuralFeels.')]
for x, label, body in steps:
    box(s, x, 2.15, 2.94, 1.24, NAVY)
    txt(s, x+.16, 2.33, 2.62, .26, label, 13, '66DDD5', True)
    txt(s, x+.16, 2.78, 2.62, .58, body, 16, WHITE)
for x in [3.47, 6.59, 9.71]: arrow(s, x, 2.69, .13)
for x, title, body in [
    (.5, 'Touch becomes an image', 'Each virtual fingertip measures contact against an untouched reference. Calibrated rendering turns this into a touch image.'),
    (4.67, 'Geometry provides a reference', 'The simulator knows the object’s shape and position. These provide a reference for evaluating an estimated 3D shape.'),
    (8.84, 'Data connects to perception', 'The exporter matches NeuralFeels, a system for estimating object shape and motion from vision and touch.')]:
    box(s, x, 3.75, 3.95, 2.16, LIGHT)
    txt(s, x+.18, 3.96, 3.59, .5, title, 17, TEAL, True)
    txt(s, x+.18, 4.59, 3.59, 1.15, body, 16)
takeaway(s, 'New research use: study how touch contributes to 3D perception during active object manipulation.')
s.notes_slide.notes_text_frame.text = '''Presenter explanation: Imagine recording a manipulation experiment with synchronized external cameras and four fingertip microscopes. In simulation, the object geometry and position are also available, providing a reference that is difficult to measure precisely in physical experiments.
The original repository did not include these collectors or the NeuralFeels exporter. New collectors can use the first-stage teacher or second-stage adapted policy. The original two-stage learning design is retained.
NeuralFeels is the downstream perception consumer. This work exports its expected inputs; it does not establish improved reconstruction accuracy. Touch is recorded for perception research, not fed into the inherited movement policy.
Known object geometry is exported as a signed distance field (SDF): a 3D grid recording distance to the object surface, with a sign distinguishing inside from outside. Default grid pitch is 0.5 mm. That is a sampling resolution, not a claim of 0.5 mm reconstruction accuracy.
The soft pad uses a compliant contact model on rigid geometry, not a full soft-body material simulation. Rendering and calibration are documented implementation choices, not proof of physical sensor fidelity.
''' + notes[1] + '\n' + notes[2]

s = slide(3, 'Development difficulties: making the pieces agree',
    'The difficult work was preserving physical meaning across simulator, sensor and dataset conventions.',
    'Difficulties and fixes are documented in implementation comments and export checks; no performance improvement is quantified here.')
for x,w,label in [(.5,3.14,'DIFFICULTY'),(3.78,4.1,'WHY IT HAPPENED'),(8.02,4.78,'IMPLEMENTED RESPONSE')]:
    box(s,x,2.04,w,.45,NAVY)
    txt(s,x+.14,2.14,w-.28,.25,label,12,'66DDD5',True)
issues=[
    ('Keeping finger commands correct', 'The new simulator lists the same joints in a different order.', 'Reorder measurements and commands so each value reaches the intended joint.'),
    ('Getting a meaningful touch signal', 'Rigid pads barely indented; the untouched reference could include contact.', 'Allow compliant contact and record the untouched reference before the first grasp.'),
    ('Making imported objects visible', 'One object-import path produced empty or invisible geometry.', 'Create simple shapes directly; convert detailed surface models through a separate path.'),
    ('Keeping images and 3D data consistent', 'Sensor labels, 3D coordinates and recording times did not agree.', 'Map conventions, check camera poses, preserve depth and exclude reset frames.'),
]
for i,(difficulty,cause,fix) in enumerate(issues):
    y=2.65+i*.85
    box(s,.5,y,12.3,.77,LIGHT if i%2==0 else 'FAFCFD')
    txt(s,.65,y+.12,2.84,.62,difficulty,16,AMBER,True)
    txt(s,3.92,y+.1,3.79,.65,cause,15,GRAY)
    txt(s,8.16,y+.1,4.46,.65,fix,15,NAVY)
takeaway(s, 'Key lesson: a plausible image or motion is insufficient—contact, geometry and timing must agree.')
s.notes_slide.notes_text_frame.text = '''Evidence and presentation scope: These are difficulties documented in the repository, not reconstructed development dates or estimates of time spent. No simulator run, hardware trial or end-to-end perception benchmark was performed while preparing these slides.
1. Joint order: hora/algo/deploy/deploy_ros2.py and deploy_ros2_two_hands.py document the simulator import changing a finger-major ordering to a joint-rank-major ordering. An explicit self-inverse 4x4 transpose permutation aligns observations, histories and outgoing actions. Without consistent ordering, values can be assigned to the wrong joints. The slide describes the compatibility problem, not a claim that a specific hardware accident occurred.
Additional control difficulty: hora/tasks/isaaclab/allegro_hand_hora_env.py:_apply_action documents that two lazy reads access the same current state. The code stores the previous physics-step joint position separately to estimate velocity for damping.
2. Tactile signal: hora/tasks/isaaclab/allegro_hand_hora_env.py:UrdfFileWithCompliantContactCfg documents near-flat contact imagery with ordinary rigid contacts. It binds a compliant PhysX material to elastomer links. scripts/dataset_collection/sensors.py:capture_initial_tactile_render documents an already-contacting grasp contaminating the nominal baseline and a later attempted object-teleport approach causing unstable contact after restoration. The current code takes a naturally untouched baseline before reset and refreshes camera poses.
Additional grasp difficulty: hora/tasks/isaaclab/allegro_hand_grasp_env.py:_setup_scene documents all-zero filtered contact forces with one sensor spanning multiple fingertip bodies in parallel environments. One object-filtered sensor per fingertip addresses this indexing limitation and replaces the original CPU-only query with GPU-compatible detection. No measured speedup is asserted.
3. Asset visibility: hora/tasks/isaaclab/allegro_hand_hora_env.py:_build_object_cfg documents empty/invisible rigid-object geometry from a URDF-to-USD payload/reference composition path. Native primitive spawners bypass it for simple shapes; MeshConverter handles detailed surface geometry. _make_mesh_materials_opaque also addresses invisible imported materials. This is a documented importer issue in the development environment, not a universal claim about all simulator versions.
4. Data consistency: scripts/dataset_collection/neuralfeels_export.py checks camera poses against NeuralFeels forward kinematics on every captured frame. legacy_tactile.py and conventions.py define finger/joint and camera conventions. feelsight_writer.py fits depth encoding to preserve measurements for the existing loader. Both collectors render each physics step and exclude done steps because the simulator has already reset. See detailed export notes below.
''' + notes[2]

prs.core_properties.title = 'Allegro In-Hand Rotation: What Changed and Why It Was Difficult'
prs.core_properties.subject = 'Accessible comparison of original and extended research capabilities'
prs.core_properties.author = 'Research Engineering'
prs.save(OUT / 'allegro_research_changes.pptx')
print(OUT / 'allegro_research_changes.pptx')
